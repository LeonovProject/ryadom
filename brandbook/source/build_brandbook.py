from pathlib import Path
import argparse
from functools import lru_cache
import json, math, shutil, zipfile
import subprocess
import tempfile
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT=Path(__file__).resolve().parents[2]
parser=argparse.ArgumentParser(description='Build the Ryadom brandbook DOCX and design assets.')
parser.add_argument('--output-dir',type=Path,default=ROOT/'brandbook',help='Output directory; defaults to brandbook in this checkout.')
parser.add_argument('--work-dir',type=Path,help='Intermediate images; defaults to a unique temporary directory.')
parser.add_argument('--fonts-dir',type=Path,help='Directory with static Open Sans and Montserrat TTF/OTF files.')
args=parser.parse_args()
OUT=args.output_dir.resolve()
QA=args.work_dir.resolve() if args.work_dir else Path(tempfile.mkdtemp(prefix='ryadom-brandbook-'))
AS=OUT/'assets'
for p in (OUT,QA,AS): p.mkdir(parents=True,exist_ok=True)
C={'red':'#D92D3A','hover':'#BB2230','pressed':'#9F1C28','tint':'#FFF1F2','white':'#FFFFFF','black':'#17191D','muted':'#5C616A','bg':'#F6F7F9','border':'#DDE0E5','control':'#858B95','blue':'#2457C5','blueTint':'#EEF3FF','error':'#A61B29','success':'#236341','warning':'#805600'}
@lru_cache(maxsize=4)
def font_path(bold=False,display=False):
    family='Montserrat' if display else 'Open Sans'
    weight='Bold' if bold else 'Regular'
    if args.fonts_dir:
        stem=('Montserrat' if display else 'OpenSans')+'-'+weight
        for ext in ('.ttf','.otf'):
            candidates=sorted(args.fonts_dir.resolve().rglob(stem+ext))
            if candidates:
                return str(candidates[0])
        raise FileNotFoundError(f'Missing {stem}.ttf or {stem}.otf in {args.fonts_dir}')
    try:
        result=subprocess.run(['fc-match','-f','%{file}\n%{family}\n',f'{family}:style={weight}'],check=True,capture_output=True,text=True)
    except (FileNotFoundError,subprocess.CalledProcessError) as exc:
        raise RuntimeError('Install Open Sans and Montserrat, or pass --fonts-dir with their static font files.') from exc
    lines=result.stdout.splitlines()
    if len(lines)<2 or family.casefold() not in lines[1].casefold() or not Path(lines[0]).is_file():
        raise RuntimeError(f'{family} {weight} is unavailable; install it or pass --fonts-dir.')
    return lines[0]

def f(size,bold=False,display=False):
    return ImageFont.truetype(font_path(bold,display),size)
def canvas(w,h,color='white'):
    im=Image.new('RGB',(w,h),color); return im,ImageDraw.Draw(im)
def txt(d,xy,s,size=28,color=None,bold=False,display=False):
    d.text(xy,s,font=f(size,bold,display),fill=color or C['black'],spacing=int(size*.3))
def rr(d,box,fill,r=18,outline=None,width=2): d.rounded_rectangle(box,r,fill=fill,outline=outline,width=width)
def mark(d,x,y,s=100,color=None):
    # 64 unit master: paired lines with identical weight and direction.
    for dx in [0,24]:
        d.line((x+(13+dx)*s/64,y+49*s/64,x+(27+dx)*s/64,y+15*s/64),fill=color or C['red'],width=round(12*s/64))
        for xx,yy in [(13+dx,49),(27+dx,15)]:
            r=6*s/64; d.ellipse((x+xx*s/64-r,y+yy*s/64-r,x+xx*s/64+r,y+yy*s/64+r),fill=color or C['red'])
def logo(d,x,y,s=100,inverse=False):
    mark(d,x,y,s,C['white'] if inverse else C['red'])
    txt(d,(x+s*1.12,y+s*.03),'рядом',round(s*.68),C['white'] if inverse else C['black'],True,True)
def button(d,x,y,w,label,kind='primary'):
    fill=C['red'] if kind=='primary' else (C['border'] if kind=='disabled' else C['white'])
    rr(d,(x,y,x+w,y+76),fill,18, C['control'] if kind=='secondary' else None)
    font=f(25,True); tw=d.textbbox((0,0),label,font=font)[2]
    txt(d,(x+(w-tw)/2,y+20),label,25,C['white'] if kind=='primary' else (C['muted'] if kind=='disabled' else C['black']),True)
def chip(d,x,y,label,active=False):
    fw=d.textbbox((0,0),label,font=f(23))[2]+44
    rr(d,(x,y,x+fw,y+50),C['tint'] if active else C['bg'],25,C['red'] if active else None)
    txt(d,(x+22,y+9),label,23,C['pressed'] if active else C['black'])
    return fw
def save(im,name): im.save(QA/(name+'.png')); return QA/(name+'.png')

# Cover plate is a brand application, separate from the document title.
im,d=canvas(1500,760)
logo(d,35,20,150)
txt(d,(35,230),'Есть планы.\nБудет компания.',76,bold=True,display=True)
txt(d,(39,480),'Спорт, прогулки и встречи по интересам',30,C['muted'])
x=40
for s in ['Бег','Прогулки','Настольные игры']:
    x+=chip(d,x,590,s,s=='Бег')+14
mark(d,1090,175,360)
cover=save(im,'cover')

im,d=canvas(1500,570)
logo(d,45,70,170)
rr(d,(810,20,1490,285),C['red'],26); logo(d,855,80,140,True)
rr(d,(810,315,1490,550),C['black'],26); logo(d,855,360,140,True)
txt(d,(48,290),'Основная версия',30,bold=True)
txt(d,(48,345),'Знак + название на светлом фоне',25,C['muted'])
mark(d,50,440,80)
txt(d,(160,445),'Знак отдельно — аватар и иконка',25)
logos=save(im,'logos')

im,d=canvas(1500,370)
swatches=[('red','Действие'),('white','Воздух'),('black','Текст'),('blue','Информация')]
widths=[520,470,300,210]; x=0
for (key,label),w in zip(swatches,widths):
    rr(d,(x,0,x+w-8,240),C[key],18,C['border'] if key=='white' else None)
    txt(d,(x+20,270),C[key],25,bold=True)
    txt(d,(x+20,310),label,23)
    x+=w
palette=save(im,'palette')

im,d=canvas(1500,680,C['bg'])
txt(d,(40,28),'Действия',28,bold=True)
button(d,40,90,390,'Присоединиться')
button(d,470,90,340,'Создать встречу','secondary')
button(d,850,90,350,'Нет мест','disabled')
txt(d,(40,220),'Поиск и выбор',28,bold=True)
rr(d,(40,285,750,365),C['white'],16,C['control'])
txt(d,(65,305),'Название или место',26,C['muted'])
chip(d,800,300,'Бег',True);chip(d,945,300,'Прогулки')
txt(d,(40,420),'Ошибка с понятным объяснением',28,bold=True)
txt(d,(40,478),'Время встречи',24,bold=True)
rr(d,(40,522,440,600),C['white'],16,C['error'])
txt(d,(65,542),'18:30',26)
txt(d,(480,540),'!  Выбери время позже текущего.',25,C['error'])
components=save(im,'components')

im,d=canvas(1500,1010,C['bg'])
rr(d,(16,16,1484,994),C['white'],24,C['border'])
logo(d,55,45,82)
txt(d,(510,65),'Встречи     Сообщества     Мои встречи',24)
button(d,1110,43,325,'Создать встречу','secondary')
txt(d,(65,185),'Найди компанию\nдля своих планов',66,bold=True,display=True)
txt(d,(69,375),'Тренировка, прогулка или вечер за игрой.',28,C['muted'])
rr(d,(65,450,1435,548),C['bg'],20)
txt(d,(93,482),'Что хочется делать?',28,C['muted'])
txt(d,(625,482),'Город или район',28,C['muted'])
button(d,1110,461,310,'Найти встречу')
x=65
for s in ['Все','Бег','Прогулки','Настольные игры']:
    x+=chip(d,x,585,s,s=='Все')+12
txt(d,(65,685),'Ближайшие встречи',34,bold=True,display=True)
for x,title,cat,time,level in [(65,'Бежим вдоль реки','Бег','Сегодня, 19:00','Спокойный темп'),(530,'Прогулка по парку','Прогулки','Завтра, 11:00','Подойдёт всем'),(995,'Вечер настольных игр','Игры','Суббота, 18:00','Объясним правила')]:
    rr(d,(x,755,x+430,955),C['white'],18,C['border'])
    txt(d,(x+22,775),cat.upper(),19,C['red'],True)
    txt(d,(x+22,817),title,25,bold=True)
    txt(d,(x+22,866),time,24)
    txt(d,(x+22,910),level,21,C['muted'])
web=save(im,'web')

im,d=canvas(1500,1110)
rr(d,(45,20,620,1090),C['black'],65)
rr(d,(60,35,605,1075),C['white'],55)
txt(d,(95,68),'9:41',22,bold=True)
logo(d,95,125,80)
txt(d,(95,240),'Планы на сегодня',33,bold=True,display=True)
rr(d,(95,310,570,390),C['bg'],16,C['control'])
txt(d,(119,334),'Что хочется делать?',25,C['muted'])
chip(d,95,424,'Бег',True);chip(d,220,424,'Прогулки')
rr(d,(95,520,570,880),C['white'],22,C['border'])
txt(d,(120,548),'Бег · Спокойный темп',23,C['red'])
txt(d,(120,599),'Бежим вдоль реки',29,bold=True)
txt(d,(120,658),'Сегодня, 19:00',27)
txt(d,(120,705),'Парк у набережной',24,C['muted'])
txt(d,(120,752),'4 из 8 участников',24,C['muted'])
button(d,120,798,425,'Присоединиться')
d.line((80,960,585,960),fill=C['border'],width=2)
txt(d,(93,990),'Встречи    Мои встречи    Профиль',21)
txt(d,(700,110),'Одна система\nна разных экранах',49,bold=True,display=True)
txt(d,(705,310),'Встреча понятна до записи:',28,bold=True)
for i,s in enumerate(['когда и где встречаемся','какой темп или формат','сколько осталось мест','кто организует встречу']):
    txt(d,(705,380+i*68),'•  '+s,27)
txt(d,(705,720),'48',95,C['red'],True,True)
txt(d,(710,845),'единиц — базовая высота кнопки',25,C['muted'])
txt(d,(710,890),'и минимальная зона касания в продукте',25,C['muted'])
mobile=save(im,'mobile')

im,d=canvas(1500,700)
rr(d,(0,0,730,700),C['red'],24)
logo(d,45,40,90,True)
txt(d,(50,200),'Собираемся\nна пробежку',59,C['white'],True,True)
txt(d,(52,410),'Суббота · 10:00\nПарк у набережной\nСпокойный темп',30,C['white'])
rr(d,(770,0,1500,700),C['bg'],24)
logo(d,815,40,90)
txt(d,(820,200),'Собираемся\nза одним столом',56,bold=True,display=True)
txt(d,(822,410),'Пятница · 19:00\nКафе в центре\nНастольные игры',30,C['muted'])
mark(d,1260,500,180)
campaign=save(im,'campaign')


# Product scenarios, shared with the visual system.
def badge(d,x,y,label,tone='neutral'):
    colors={'neutral':(C['bg'],C['muted']),'info':(C['blueTint'],C['blue']),'success':('#EDF6F0',C['success']),'warning':('#FFF7E7',C['warning']),'error':(C['tint'],C['error'])}
    bg,fg=colors[tone]; w=d.textbbox((0,0),label,font=f(24,True))[2]+68
    rr(d,(x,y,x+w,y+56),bg,28)
    if tone=='success':
        d.line((x+17,y+29,x+24,y+36,x+36,y+20),fill=fg,width=3)
    elif tone=='info':
        d.ellipse((x+16,y+17,x+38,y+39),outline=fg,width=2)
        d.line((x+27,y+21,x+27,y+28,x+32,y+31),fill=fg,width=2)
    elif tone=='error':
        d.line((x+18,y+19,x+35,y+36),fill=fg,width=3)
        d.line((x+35,y+19,x+18,y+36),fill=fg,width=3)
    else:
        d.ellipse((x+21,y+23,x+31,y+33),fill=fg)
    txt(d,(x+48,y+11),label,24,fg,True)
    return w

im,d=canvas(1500,310)
txt(d,(30,20),'Статус дополняет действие',33,bold=True,display=True)
badge(d,30,100,'Черновик');badge(d,330,100,'На проверке','info');badge(d,710,100,'Ты участвуешь','success')
badge(d,30,200,'Нужны уточнения','warning');badge(d,520,200,'Заявка отклонена','error')
stateplate=save(im,'states')

im,d=canvas(1500,590,C['bg'])
for x,title,mode,action,detail in [(30,'Бежим вдоль реки','Свободное участие','Присоединиться','После присоединения'),(780,'Играем в волейбол','По заявке','Подать заявку','После отправки заявки')]:
    rr(d,(x,25,x+690,565),C['white'],22,C['border'])
    txt(d,(x+30,58),mode.upper(),23,C['red'],True)
    txt(d,(x+30,115),title,34,bold=True,display=True)
    txt(d,(x+30,195),'Сегодня, 19:00 · 4 из 8 участников',25)
    txt(d,(x+30,246),'Организатор: клуб «Темп»',25,C['muted'])
    button(d,x+30,320,630,action)
    txt(d,(x+32,427),detail,25,C['muted'])
    if x==30: badge(d,x+30,488,'Ты участвуешь','success')
    else: badge(d,x+30,488,'Заявка отправлена','info')
joinplate=save(im,'participation')

im,d=canvas(1500,510,C['bg'])
rr(d,(25,25,750,485),C['white'],20,C['border'])
txt(d,(55,55),'Новая встреча',37,bold=True,display=True)
txt(d,(55,130),'1  О встрече     2  Участие     3  Проверка',24,C['muted'])
txt(d,(55,207),'Как можно присоединиться?',29,bold=True)
chip(d,55,268,'Свободно',True);chip(d,265,268,'По заявке')
button(d,55,365,350,'Сохранить черновик','secondary')
rr(d,(790,25,1475,485),C['white'],20,C['border'])
txt(d,(820,55),'Организация',37,bold=True,display=True)
badge(d,820,145,'На проверке','info')
txt(d,(820,235),'Публикация станет доступна\nпосле допуска организации.',29)
txt(d,(820,340),'Черновик можно сохранить.',25,C['muted'])
publishplate=save(im,'publication')

im,d=canvas(1500,570,C['bg'])
rr(d,(25,25,725,545),C['white'],22,C['border'])
txt(d,(55,55),'Уточнить условия',36,bold=True,display=True)
txt(d,(55,135),'Нужна ли своя ракетка?',29,bold=True)
txt(d,(55,200),'Организатор',23,C['muted'])
txt(d,(55,252),'Можно прийти без неё.\nИнвентарь выдадим на месте.',28)
button(d,55,420,400,'Задать вопрос','secondary')
rr(d,(770,25,1475,545),C['white'],22,C['border'])
txt(d,(805,58),'Снять с публикации?',35,bold=True,display=True)
txt(d,(805,155),'Встреча перестанет отображаться\nв каталоге.',28)
txt(d,(805,270),'Это действие управляет публикацией.',23,C['muted'])
button(d,805,350,620,'Снять с публикации')
txt(d,(1005,462),'Оставить в каталоге',25,bold=True)
manageplate=save(im,'clarify-and-unpublish')

im,d=canvas(1500,550)
logo(d,30,15,95)
txt(d,(30,155),'Те же планы.\nВсегда под рукой.',55,bold=True,display=True)
txt(d,(35,345),'Встречи   /   Мои встречи   /   Профиль',25,C['muted'])
rr(d,(790,35,1470,250),C['bg'],24,C['border'])
txt(d,(820,65),'Рядом · Заявки',23,C['muted'])
txt(d,(820,115),'Заявка одобрена',31,bold=True)
txt(d,(820,176),'Открой встречу, чтобы посмотреть детали.',24)
rr(d,(790,295,1470,510),C['bg'],24,C['border'])
txt(d,(820,325),'Рядом · Встречи',23,C['muted'])
txt(d,(820,375),'Изменилось место встречи',29,bold=True)
txt(d,(820,436),'Посмотри новый адрес перед выходом.',24)
nativeplate=save(im,'mobile-extension')

# Portable logo masters. Text stays editable; mark has no font dependency.
symbol='<path d="M13 49 L27 15 M37 49 L51 15" fill="none" stroke="{color}" stroke-width="12" stroke-linecap="round"/>'
for name,col,bg in [('mark',C['red'],''),('mark-white',C['white'],''),('mark-black',C['black'],'')]:
    (AS/(name+'.svg')).write_text('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 64 64" role="img" aria-label="Рядом">'+symbol.format(color=col)+'</svg>')
for name,col,textcol,bg in [('logo',C['red'],C['black'],''),('logo-inverse',C['white'],C['white'],C['red']),('logo-mono',C['black'],C['black'],'')]:
    b=f'<rect width="252" height="80" rx="12" fill="{bg}"/>' if bg else ''
    svg=f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 252 80" role="img" aria-label="Рядом">{b}<g transform="translate(8 8)">{symbol.format(color=col)}</g><text x="80" y="53" font-family="Montserrat, sans-serif" font-weight="700" font-size="44" fill="{textcol}">рядом</text></svg>'
    (AS/(name+'.svg')).write_text(svg)
(AS/'app-icon.svg').write_text(f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1024 1024"><rect width="1024" height="1024" fill="{C["red"]}"/><g transform="translate(176 176) scale(10.5)">{symbol.format(color=C["white"])}</g></svg>')

tokens={'color':{'brand':{'primary':C['red'],'hover':C['hover'],'pressed':C['pressed'],'tint':C['tint'],'onTint':C['pressed']},'surface':{'base':C['white'],'subtle':C['bg']},'text':{'primary':C['black'],'secondary':C['muted'],'inverse':C['white']},'border':{'subtle':C['border'],'control':C['control']},'info':{'foreground':C['blue'],'background':C['blueTint']},'status':{'error':C['error'],'success':C['success'],'warning':C['warning']},'focus':C['blue']},'space':{'1':4,'2':8,'3':12,'4':16,'6':24,'8':32,'12':48,'16':64},'radius':{'control':12,'card':16,'dialog':24,'pill':999},'size':{'control':48,'icon':24,'contentMax':1200},'font':{'body':'Open Sans','heading':'Montserrat','weights':{'regular':400,'semibold':600,'bold':700}},'motion':{'fast':120,'normal':180,'slow':240},'breakpoint':{'compact':600,'wide':1024}}

tokens['color']['statusBackground']={'neutral':C['bg'],'info':C['blueTint'],'success':'#EDF6F0','warning':'#FFF7E7','error':C['tint']}
tokens['type']={
 'web':{'display':{'fontSize':48,'lineHeight':56},'pageTitle':{'fontSize':32,'lineHeight':40},'sectionTitle':{'fontSize':24,'lineHeight':32}},
 'native':{'display':{'fontSize':32,'lineHeight':40},'pageTitle':{'fontSize':28,'lineHeight':36},'sectionTitle':{'fontSize':22,'lineHeight':30}},
 'shared':{'cardTitle':{'fontSize':18,'lineHeight':26},'body':{'fontSize':16,'lineHeight':24},'caption':{'fontSize':14,'lineHeight':20},'button':{'fontSize':16,'lineHeight':24}}
}
states={
 'organization':{'draft':{'label':'Заявка не отправлена','tone':'neutral'},'pending':{'label':'На проверке','tone':'info'},'approved':{'label':'Допущена к публикации','tone':'success'},'needsInformation':{'label':'Нужны уточнения','tone':'warning'},'rejected':{'label':'Допуск не предоставлен','tone':'error'}},
 'publication':{'draft':{'label':'Черновик','tone':'neutral'},'published':{'label':'Опубликовано','tone':'success'},'unpublished':{'label':'Снято с публикации','tone':'neutral'}},
 'participation':{'pending':{'label':'Заявка отправлена','tone':'info'},'confirmed':{'label':'Ты участвуешь','tone':'success'},'rejected':{'label':'Заявка отклонена','tone':'error'}}
}
copy={
 'activity':{'create':'Создать встречу','saveDraft':'Сохранить черновик','preview':'Предпросмотр','publish':'Опубликовать','saveSuccess':'Черновик сохранён. Можно вернуться к нему позже.','publishSuccess':'Встреча опубликована. Она доступна в каталоге.','saveError':'Не удалось сохранить встречу. Попробуй ещё раз.','publishError':'Не удалось опубликовать встречу. Попробуй ещё раз.'},
 'organization':{'submit':'Отправить на проверку','pending':'Организация на проверке. Решение появится в кабинете.','approved':'Организация допущена к публикации.','needsInformation':'Для рассмотрения нужны уточнения. Посмотри комментарий в кабинете.','rejected':'Допуск не предоставлен. Посмотри причину в кабинете.','publishBlocked':'Публикация станет доступна после допуска организации.'},
 'participation':{'join':'Присоединиться','apply':'Подать заявку','pending':'Заявка отправлена. Ожидаем решение организатора.','confirmed':'Ты в списке участников. Детали есть на странице встречи.','rejected':'Заявка отклонена. Можно выбрать другую встречу.','withdraw':'Отозвать заявку','joinError':'Не удалось присоединиться. Проверь статус участия перед повторной попыткой.','applyError':'Не удалось отправить заявку. Проверь её статус перед повторной попыткой.'},
 'questions':{'open':'Уточнить условия','send':'Задать вопрос','sent':'Вопрос отправлен. Ответ появится в этом разделе.','failed':'Не удалось отправить вопрос. Попробуй ещё раз.'},
 'publication':{'unpublish':'Снять с публикации','confirmTitle':'Снять встречу с публикации?','confirmBody':'Встреча перестанет отображаться в каталоге.','keep':'Оставить в каталоге','success':'Встреча снята с публикации. Она больше не отображается в каталоге.','error':'Не удалось снять встречу с публикации. Попробуй ещё раз.'}
}
(AS/'states.json').write_text(json.dumps(states,ensure_ascii=False,indent=2))
(AS/'copy.ru.json').write_text(json.dumps(copy,ensure_ascii=False,indent=2))

(AS/'tokens.json').write_text(json.dumps(tokens,ensure_ascii=False,indent=2))
flat={}
def flatten(obj,prefix=''):
    for key,v in obj.items():
        k=(prefix+'-'+key).strip('-')
        if isinstance(v,dict): flatten(v,k)
        else: flat[k]=v
flatten(tokens)
css=['/* Рядом 1.1. Load Montserrat 600/700 and Open Sans 400/600/700 separately. */',':root {']
for k,v in flat.items():
    suffix='ms' if k.startswith('motion-') else ('px' if isinstance(v,int) and not k.startswith('font-weights') else '')
    css.append(f'  --{k}: {json.dumps(v) if k in ["font-body","font-heading"] else v}{suffix};')
css+=['}','', '.ryadom-focus:focus-visible { outline: 3px solid var(--color-focus); outline-offset: 3px; }','@media (prefers-reduced-motion: reduce) {', '  .ryadom-motion { transition: none; animation: none; }','}']
(AS/'tokens.css').write_text('\n'.join(css))
(AS/'tokens.ts').write_text('// Shared values for React and React Native. Load font faces per platform.\nexport const brand = '+json.dumps(tokens,ensure_ascii=False,indent=2)+' as const;\n')
(AS/'README.md').write_text('''# Рядом — материалы бренда 1.1

Рабочее название и концепция для платформы совместных активностей.

- `mark*.svg` — векторный знак без зависимости от шрифтов.
- `logo*.svg` — редактируемая связка знака и названия. Для точного вида нужен Montserrat Bold. Перед передачей в печать преобразовать текст в кривые в векторном редакторе.
- `app-icon.svg` — квадратный мастер 1024 × 1024 без заранее скруглённых углов. Платформенные экспорты и safe areas готовятся отдельно.
- `tokens.json` — исходные значения без платформенных единиц.
- `tokens.css` — переменные для React web. Шрифты подключаются отдельно.
- `tokens.ts` — те же значения для React / React Native. Это данные, не готовые компоненты.
- `states.json` — предложенные подписи состояний по отдельным сущностям и визуальные роли; это не схема серверных переходов.
- `copy.ru.json` — русский словарь текстов для шести сценариев. Применять после проверки соответствия фактическому поведению.

Мобильное приложение — собственная идея команды, развивающая платформу. Обязательность приложения не выводится из критериев. Допуск организации, публикация встречи и участие человека — разные состояния. Оттенки состояния сопровождаются иконкой и подписью. Уточнение условий показано как вариант с вопросами на странице встречи. Возврат публикации, уведомления и доступ по прямой ссылке требуют отдельных продуктовых правил.

Основной текст — Open Sans Regular/Semibold, заголовки — Montserrat Semibold/Bold. Для React Native нужны зарегистрированные имена начертаний; не полагаться на синтетический bold. Размеры JSON трактуются как CSS px в web и логические единицы в native; типографика масштабируется настройками пользователя. JSON — источник значений, CSS и TS — его снимки. При изменении синхронизировать все три файла.

Шрифты в архив не включены. Официальные источники и лицензии:
- https://github.com/JulietaUla/Montserrat
- https://github.com/googlefonts/opensans

Название и знак — рабочая концепция; проверка товарного знака и домена не проводилась. SVG использует доступную подпись; в приложении доступное имя задаётся компонентом с учётом контекста.
''')

# Brand book, deliberately paginated with editable text.
doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.65);sec.bottom_margin=Cm(1.5);sec.left_margin=Cm(1.75);sec.right_margin=Cm(1.75)
sec.footer_distance=Cm(.7)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    s=doc.styles[name];s.font.name='Open Sans';s.font.color.rgb=RGBColor.from_string('000000')
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Open Sans')
normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.18;normal.paragraph_format.space_after=Pt(8)
for name,size in [('Title',30),('Heading 1',23),('Heading 2',13)]:
    s=doc.styles[name];s.font.name='Montserrat';s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.space_before=Pt(13);s.paragraph_format.space_after=Pt(9)
doc.styles['Subtitle'].font.size=Pt(11)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('РЯДОМ   /   БРЕНДБУК 1.1   /   ');r.font.name='Open Sans';r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('5C616A')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='Брендбук платформы Рядом';doc.core_properties.subject='Визуальная и вербальная система платформы совместных активностей';doc.core_properties.author='Рядом'

def p(s='',bold=False):
    pp=doc.add_paragraph(); rr=pp.add_run(s);rr.bold=bold;return pp
def h(s): doc.add_heading(s,2)
def page(n,title,intro):
    pp=doc.add_paragraph(f'{n:02d}   РЯДОМ');pp.paragraph_format.page_break_before=True;pp.runs[0].font.size=Pt(9)
    doc.add_heading(title,1);p(intro)
def pic(path,width=17.5):
    pp=doc.add_paragraph();pp.paragraph_format.space_after=Pt(8);r=pp.add_run();r.add_picture(str(path),width=Cm(width))
def caption(s):
    pp=p(s);pp.paragraph_format.space_after=Pt(9)
    for r in pp.runs:r.font.size=Pt(8.5);r.font.color.rgb=RGBColor.from_string('5C616A')
def table(headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for col,w in zip(t.columns,widths):col.width=Cm(w)
    for cell,text in zip(t.rows[0].cells,headers):cell.text=text
    for row in rows:
        for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
    for ri,row in enumerate(t.rows):
        for ci,cell in enumerate(row.cells):
            if widths:cell.width=Cm(widths[ci])
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'25282D' if ri==0 else ('F6F7F9' if ri%2==0 else 'FFFFFF'));pr.append(sh)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            pr.append(borders);m=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'95');el.set(qn('w:type'),'dxa');m.append(el)
            pr.append(m)
            for pp in cell.paragraphs:
                pp.paragraph_format.space_after=Pt(1);pp.paragraph_format.space_before=Pt(1);pp.paragraph_format.line_spacing=1.18
                for rr in pp.runs:
                    rr.font.size=Pt(9);rr.bold=ri==0;rr.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '17191D')
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    p().paragraph_format.space_after=Pt(0)
    return t

doc.add_heading('Брендбук платформы Рядом',0)
doc.add_paragraph('Визуальная система и принципы общения',style='Subtitle')
caption('Версия 1.1  •  19 октября 2026')
pic(cover)
p('Рядом помогает находить людей для совместных занятий и встреч. Спорт становится первым направлением, а общая система поддерживает прогулки, игры, обучение и другие интересы.')
p('Идея бренда — сделать первый шаг к совместному занятию простым и понятным. Красный приглашает действовать, белый оставляет пространство людям и событиям.',True)
h('Как пользоваться брендбуком')
p('Разделы 1–3 задают позиционирование, язык и знак. Разделы 4–9 описывают цвета, типографику, интерфейс и применение. Разделы 10–15 раскрывают роли, статусы и шесть сценариев. Раздел 16 связывает систему с React и React Native.')
caption('Рабочее название «Рядом». Формулировки и изображения экранов описывают концепцию будущего продукта.')

page(1,'Основа бренда','Платформа объединяет людей вокруг конкретного плана: что делать, где встретиться и с кем провести время.')
h('Позиционирование')
p('Рядом — платформа, на которой можно найти или собрать компанию для совместной активности. Понятная встреча с подходящим форматом помогает перейти от намерения к участию.')
h('Обещание')
p('Помогаем найти компанию и договориться о совместных планах. До записи человек видит время, место, формат, стоимость и условия участия, если они есть.')
h('Для кого создаём')
p('Участнику нужна подходящая компания и уверенность, что его уровень или интерес совпадает с форматом. Организатору нужны понятное приглашение, состав группы и удобная связь с участниками. Сообществу нужна последовательность встреч и узнаваемость.')
table(['Принцип','Как он проявляется'],[
('Открытость','Пишем, подходит ли встреча новичкам; объясняем правила и темп.'),
('Самостоятельный выбор','Даём информацию для решения; не давим рекордами, сериями и чувством вины.'),
('Конкретность','Показываем условия до записи; изменения сообщаем явно.'),
('Уважение','Не оцениваем тело, возраст, уровень подготовки или круг общения.')],[4,13.5])
h('Единый бренд для разных интересов')
p('Используем «Рядом» и названия категорий: Бег, Прогулки, Игры, Обучение. Категории различаем текстом и иконкой, сохраняя общие цвета и компоненты. Отдельные логотипы для каждого направления не нужны.')

page(2,'Язык и характер','Говорим как внимательный организатор: спокойно, по делу и с понятным приглашением.')
h('Тон общения')
p('Обращаемся на «ты», если пользовательский контекст не требует формального обращения. Используем короткие предложения и обычные слова. В приглашениях допустима лёгкая энергия; в ошибках, отменах и оплате — точность и спокойствие.')
table(['Ситуация','Формулировка'],[
('Первый экран','Найди компанию для своих планов.'),
('Приглашение','Бежим в спокойном темпе. Присоединяйся, даже если начинаешь.'),
('Пустая выдача','По этим условиям встреч пока нет. Измени фильтры или создай свою.'),
('Успешная запись','Ты в списке участников. Встречаемся сегодня в 19:00.'),
('Встреча отменена','Организатор отменил встречу. Посмотри другие планы на этот день.'),
('Ошибка записи','Не удалось записаться. Попробуй ещё раз.'),
('Уведомление','Завтра в 10:00 — прогулка по парку. Место встречи есть в карточке.')],[4.1,13.4])
h('Сообщения бренда')
p('Основное приглашение: «Есть планы. Будет компания». Описание продукта: «Спорт, прогулки и встречи по интересам». Для кнопок используем конкретные действия: «Найти встречу», «Присоединиться», «Создать встречу».')
h('Что исключаем')
p('Не пишем «Хватит лениться», «Стань лучшей версией себя» и «Все уже тренируются». Не обещаем идеальных знакомств, абсолютной безопасности или результата тренировки. Не используем слово «тусовка» как универсальное описание всех активностей.')

page(3,'Логотип и знак','Знак — две одинаковые линии, направленные вверх и расположенные рядом. Это образ общего движения без привязки к конкретному виду спорта.')
pic(logos)
h('Основная композиция')
p('Красный знак слева, название «рядом» в Montserrat Bold справа. В интерфейсе предпочтительна горизонтальная версия. Отдельный знак используем там, где название уже видно или место ограничено: иконка приложения, аватар, компактная панель.')
table(['Параметр','Правило'],[
('Мастер знака','Поле 64 × 64; толщина линий 12; наклон примерно 22° от вертикали.'),
('Охранное поле','Не менее 12 единиц мастер-сетки со всех сторон от видимых границ композиции.'),
('Минимальный размер','Знак — 24 px или 6 мм; связка с названием — 120 px или 30 мм по ширине.'),
('Фон','Белый или очень светлый; белая версия на красном и тёмном фоне.'),
('Монохром','Все элементы чёрные либо все белые; без полутонов и прозрачности.')],[4.3,13.2])
p('Не растягиваем знак, не меняем расстояние между линиями, не наклоняем название и не добавляем тени, обводки или градиенты. На фотографии используем свободный однотонный участок либо сплошную подложку.')
caption('В комплекте SVG с редактируемым названием. Для печати текст логотипа переводят в кривые после установки Montserrat Bold.')

page(4,'Цветовая система','Белый — основа, красный — главный акцент, почти чёрный — текст. Синий поддерживает информацию и фокус клавиатуры.')
pic(palette)
table(['Роль','HEX','Применение'],[
('Красный','#D92D3A','Главная кнопка, знак, выбранная категория'),
('Белый','#FFFFFF','Основной фон и поверхности'),
('Почти чёрный','#17191D','Основной текст и иконки'),
('Синий','#2457C5','Ссылки, фокус, информационные детали'),
('Вторичный текст','#5C616A','Подписи и дополнительная информация'),
('Светлый фон','#F6F7F9','Подложка разделов и спокойные поверхности'),
('Линии / поля','#DDE0E5 / #858B95','Декоративные разделители / границы полей'),
('Мягкие акценты','#FFF1F2 / #EEF3FF','Красные / синие подложки')],[4.3,4,9.2])
h('Баланс')
p('Ориентир для обычного экрана: 80% белого и светлого, 12% красного, 6% тёмного, до 2% синего. Это композиционный ориентир, а не квота. Для рекламного носителя допустим сплошной красный фон.')
h('Контраст')
p('Белый на красном — 4,79:1; красный на белом — 4,79:1; основной текст на белом — 17,60:1; вторичный — 6,23:1; синий — 6,47:1. Значения рассчитаны для сплошных цветов. Для обычного текста целимся минимум в 4,5:1 [1].')
caption('Красный hover #BB2230; pressed #9F1C28. Синий не используется как второй основной цвет кнопки. Сочетание красного и синего текста друг на друге исключаем.')

page(5,'Шрифты и композиция','Montserrat задаёт характер названия и крупных заголовков. Open Sans обеспечивает спокойный ритм интерфейса и длинного текста [2, 3].')
sample=p('Есть планы.\nБудет компания.',True).runs[0];sample.font.size=Pt(25);sample.font.name='Montserrat'
caption('Пример рекламной формулировки. В продукте заголовок описывает действие или содержание экрана.')
table(['Стиль','Веб px / интерлиньяж','Мобильный / интерлиньяж','Шрифт'],[
('Крупный заголовок','48 / 56','32 / 40','Montserrat 700'),
('Заголовок страницы','32 / 40','28 / 36','Montserrat 700'),
('Заголовок раздела','24 / 32','22 / 30','Montserrat 600'),
('Заголовок карточки','18 / 26','18 / 26','Open Sans 600'),
('Основной текст','16 / 24','16 / 24','Open Sans 400'),
('Подпись','14 / 20','14 / 20','Open Sans 400'),
('Кнопка','16 / 24','16 / 24','Open Sans 600')],[4,4.2,4.8,4.5])
h('Сетка и отступы')
p('Базовый шаг — 4. Основные интервалы: 8, 12, 16, 24, 32, 48 и 64. На широком экране используем 12 колонок, промежуток 24 и контейнер до 1200 px. На мобильном — 4 колонки, поля 16 и промежуток 16; на планшете — 8 колонок и поля 24.')
h('Форма')
p('Скругления: 12 у полей и кнопок, 16 у карточек, 24 у диалогов. Полное скругление оставляем фильтрам и аватарам. Высота основных контролов — 48. Тени используем только для плавающих элементов; карточкам достаточно линии или изменения фона.')
h('Читаемость')
p('Текст выравниваем по левому краю. Не набираем длинные фразы прописными буквами. Ограничиваем строку примерно 60–75 символами. На мобильном разрешаем системное увеличение текста и рост компонентов по содержимому.')

page(6,'Компоненты интерфейса','На экране выделяем одно главное действие. Остальные действия поддерживают его и сохраняют ясную иерархию.')
pic(components,16)
h('Кнопки и состояния')
p('Основная кнопка красная с белым текстом. Вторичная — светлая с тёмным текстом и заметной границей. Высота от 48, горизонтальные отступы от 20. Hover затемняет красный; нажатие затемняет сильнее; загрузка сохраняет ширину и показывает «Записываем…». Повторное нажатие во время отправки блокируется.')
p('Состояние «нет мест» использует нейтральную заливку и явную подпись. После успешной записи показываем «Ты участвуешь» и отдельное действие отмены. Ошибка не заменяет текст кнопки без объяснения рядом.')
h('Карточка встречи')
p('Обязательные поля: название, категория, время и часовой пояс при необходимости, место, темп или формат, участники и вместимость, организатор, стоимость либо «Бесплатно». Если условия неизвестны, показываем это явно. Для неспортивных встреч «темп» заменяется подходящим описанием формата.')
h('Поля и статусы')
p('Подпись поля остаётся видимой при вводе. Фокус обозначается синим контуром 3 px с отступом 3 px. Ошибка — тёмно-красный #A61B29, значок и пояснение; успех — #236341 с подписью; предупреждение — #805600 на светлой подложке. Цвет никогда не передаёт состояние один [1].')
caption('Светлая линия #DDE0E5 используется как декор; граница поля #858B95 помогает различать контрол на белом фоне. Схематичный набор компонентов не является готовой библиотекой.')

page(7,'Пример веб интерфейса','Поиск начинается с желания человека, а карточки дают достаточно информации для выбора следующего шага.')
pic(web,15.5)
caption('Демонстрационный макет главной страницы. Названия встреч и количество участников приведены для примера.')
h('Композиция')
p('Верхняя панель содержит бренд, разделы и создание встречи. В первом экране — приглашение, поиск по активности и месту, затем категории и ближайшие события. Красный выделяет поиск и выбранные элементы; светлый фон удерживает визуальный баланс.')
h('Поведение на узком экране')
p('Колонки карточек сокращаются с трёх до двух и одной. Поля поиска располагаются вертикально. Категории переносятся или прокручиваются с видимым продолжением. Нижняя мобильная навигация содержит основные разделы; создание встречи доступно отдельным действием.')
p('Путь участника: выбор встречи, подробности и условия, присоединение, подтверждение. Дополнительные функции используют ту же визуальную систему.')

page(8,'Пример мобильного приложения','Мобильное приложение — собственная идея команды для развития Рядом. Оно сохраняет знакомые цвета, тексты и порядок информации веба.')
pic(mobile)
caption('Демонстрационная композиция для будущего приложения. Подробности организатора и условия раскрываются на экране встречи.')
h('Правила адаптации')
p('Веб и приложение используют одни роли цветов и одну шкалу отступов. В React Native размеры применяются как логические единицы. Учитываем safe area, системные жесты и увеличение шрифта. Высота кнопки 48 — минимум; при крупном тексте она увеличивается.')
h('Доступность')
p('Сохраняем понятные подписи навигации, доступные имена и порядок чтения. Элементы доступны с клавиатуры на вебе. Видимый фокус не перекрывается закреплёнными панелями. Для интерактивных целей целимся в область от 48 × 48; для веба отдельно проверяем WCAG 2.2 [1].')

page(9,'Графика и коммуникации','Узнаваемость строится на паре направленных линий, крупной типографике, свободном пространстве и конкретном приглашении.')
pic(campaign)
caption('Два примера носителей показывают, как одна система работает для спорта и других совместных занятий.')
h('Фирменная графика')
p('Парные линии можно увеличивать и обрезать краем носителя, сохраняя их толщину, наклон и взаимное положение. На одном носителе достаточно одного крупного жеста. В интерфейсе графика помогает пустым состояниям и обложкам, не мешая чтению.')
h('Фотографии и иллюстрации')
p('Показываем людей в процессе совместного занятия: общий маршрут, разговор перед стартом, игру за столом. В кадре — разный опыт и естественные эмоции. Свет нейтральный или дневной, кожа естественного цвета. Красный можно поддержать деталью одежды или предметом.')
p('Избегаем постановочной агрессии, культа идеального тела, тяжёлых цветных фильтров и ощущения закрытого клуба. Иллюстрации простые и округлые; они изображают людей и ситуацию, а не только спортивный инвентарь.')
h('Движение')
p('Переходы длятся 120–240 мс: появление, лёгкое смещение, изменение состояния. Не используем постоянное подпрыгивание кнопок и мигающие акценты. При настройке уменьшения движения отключаем декоративную анимацию.')


page(10,'Роли и доверие','Встреча объединяет участников, её автора и, при необходимости, организацию. Бренд сохраняет один характер во всех кабинетах.')
h('Роли в интерфейсе')
table(['Контекст','Что важно человеку','Язык действий'],[
('Участник','Условия встречи и результат участия','Присоединиться, подать заявку, уточнить условия'),
('Организатор','Содержание встречи, состав и заявки','Создать встречу, посмотреть заявки, снять с публикации'),
('Организация','Статус допуска и публикация событий','Отправить на проверку, посмотреть решение'),
('Проверяющий','Основания решения и следующий шаг','Предоставить допуск, запросить уточнения')],[3.3,6.3,7.9])
h('Роль зависит от контекста')
p('Один человек может участвовать в одной встрече и организовывать другую. Показываем доступные действия в контексте встречи или организации. Не создаём отдельные визуальные бренды для каждого типа кабинета.')
h('Что означает допуск')
p('Подпись «Допущена к публикации» сообщает право организации публиковать события. Она не означает гарантию качества каждой встречи. Основания допуска раскрываются в пояснении, когда они определены продуктовой политикой.')
h('Как показываем организатора')
p('В карточке и на странице встречи указываем автора или организацию. Название и аватар помогают узнавать организатора. Статус допуска показываем рядом с организацией, а статус участия — рядом с действием участника.')
h('Спокойный рабочий кабинет')
p('В формах проверки и управления используем белый фон, тёмный текст и компактные статусы. Красный оставляем главному действию. В решении об отказе объясняем причину и допустимый следующий шаг; не обещаем срок ответа без установленного процесса.')
caption('Роли и конкретные действия кабинетов — предлагаемая модель для проектирования. Её связывают с правилами допуска и публикации при детализации продукта.')

page(11,'Статусы и сообщения','Действие отвечает на вопрос «что сделать», статус — «что происходит сейчас». Цвет поддерживает смысл, подпись делает его однозначным.')
pic(stateplate)
table(['Сущность','Подписи состояний','Визуальная роль'],[
('Организация','На проверке / Допущена к публикации','Информация / успех'),
('Организация','Нужны уточнения / Допуск не предоставлен','Предупреждение / отказ'),
('Встреча','Черновик / Опубликовано / Снято с публикации','Нейтральный / успех / нейтральный'),
('Участие','Заявка отправлена / Ты участвуешь','Информация / успех'),
('Участие','Заявка отклонена','Отказ с пояснением')],[3.2,8.7,5.6])
h('Цвета статусов')
p('Нейтральный: #5C616A на #F6F7F9. Информация: #2457C5 на #EEF3FF. Успех: #236341 на #EDF6F0. Уточнение: #805600 на #FFF7E7. Отказ и ошибка: #A61B29 на #FFF1F2. Используем короткую подпись и простую иконку; размер иконки 16–20, основной набор — 24.')
h('Разделяем разные процессы')
p('Организация может быть допущена, встреча ещё оставаться черновиком, а заявка человека ожидать решения. Поэтому один общий бейдж «Одобрено» не заменяет три независимых состояния.')
h('Пишем результат и следующий шаг')
p('«Заявка отправлена. Ожидаем решение организатора». «Допуск не предоставлен. Посмотри причину в кабинете». «Встреча снята с публикации. Она больше не отображается в каталоге». Не используем «Готово» без указания результата.')
caption('Имена состояний в states.json относятся к визуальной системе. Серверные переходы и права задаются отдельно. Набор допуска включает вариант запроса уточнений.')

page(12,'Создание и публикация','Форма помогает сформулировать понятное приглашение, а публикация делает встречу доступной в каталоге.')
pic(publishplate,16.5)
h('Создание встречи')
p('Предлагаем три шага: «О встрече», «Участие», «Проверка». Сначала задаются название, категория, время, место и описание. Затем — способ присоединения и условия. Предпросмотр показывает ту же структуру, которую увидит участник.')
h('Подписи действий')
p('«Сохранить черновик» сохраняет незавершённую встречу. «Предпросмотр» помогает проверить приглашение. «Опубликовать» используется, когда условия публикации выполнены. После действия подтверждаем конкретный результат: «Черновик сохранён» или «Встреча опубликована».')
h('Допуск организации')
p('До отправки объясняем, какие данные нужны. После отправки показываем «На проверке» и место появления решения. Если публикация зависит от допуска, объяснение находится рядом с действием. Возможность сохранения черновика в таком состоянии — предлагаемый вариант поведения.')
h('Ошибки и восстановление')
p('Показываем ошибку у соответствующего поля и сохраняем введённое. При проблеме отправки оставляем форму доступной для повторной попытки. Не показываем «Опубликовано», пока результат не подтверждён системой.')
caption('Сценарии 1 и 2. Состав полей и правила публикации обычным пользователем уточняются в продуктовой спецификации; визуальная система поддерживает выбранные правила.')

page(13,'Присоединение и заявки','До нажатия человек понимает, вступает ли он сразу или ожидает решения организатора.')
pic(joinplate,16.5)
h('Свободное участие')
p('Основная кнопка — «Присоединиться». После подтверждения показываем «Ты участвуешь». Если мест нет, явно сообщаем «Нет мест». Во время отправки сохраняем размер кнопки и не допускаем повторного действия.')
h('Участие по заявке')
p('Основная кнопка — «Подать заявку». После отправки показываем «Заявка отправлена. Ожидаем решение организатора». Ожидание не выглядит как подтверждённое участие. Предлагаем действие «Отозвать заявку», если оно предусмотрено правилами продукта.')
h('Решение организатора')
p('В кабинете показываем заявки отдельно от участников. Действия «Принять» и «Отклонить» сопровождаем результатом. После одобрения участник видит «Ты участвуешь». При отказе используем спокойную формулировку, без оценки человека.')
h('Общие правила')
p('Состояние участия видно при повторном открытии встречи. При неопределённом результате отправки сначала проверяется актуальный статус. Текст ошибки предлагает следующий шаг, а не приглашает бесконечно нажимать ту же кнопку.')
caption('Сценарии 3 и 4. Макет демонстрирует различие способов участия; подтверждение, отказ и отзыв должны соответствовать согласованным продуктовым правилам.')

page(14,'Уточнение и управление','Вопрос об условиях помогает принять решение. Управление публикацией требует ясного описания последствий.')
pic(manageplate,16.5)
h('Уточнение условий')
p('Предлагаемый вариант — раздел вопросов на странице встречи. Действие «Уточнить условия» открывает контекст встречи, «Задать вопрос» отправляет текст. Показываем автора ответа и сохраняем связь с конкретной встречей. Отправку вопроса не обозначаем как получение ответа.')
p('Общие вопросы можно показывать всем, если это поддерживает выбранная модель общения. Личные данные и приватные вопросы требуют подходящего отдельного канала. Если выбран чат, используем те же термины и подписи результатов.')
h('Снятие с публикации')
p('Действие называется «Снять с публикации». Перед подтверждением объясняем, что встреча исчезнет из каталога. Альтернатива — «Оставить в каталоге». После завершения показываем статус «Снято с публикации».')
h('Разные последствия требуют разных слов')
p('Снятие с публикации, отмена встречи и удаление записи — отдельные действия. Доступ по прямой ссылке, уведомления участникам и возможность вернуть публикацию задаются продуктовой политикой. Существенные последствия сообщаются до подтверждения действия.')
caption('Сценарии 5 и 6. Блок вопросов — дизайн-вариант, а не закреплённая критерием механика. На макете снятие управляет видимостью в каталоге.')

page(15,'Развитие мобильного приложения','Мобильное приложение — собственная продуктовая идея команды. Оно продолжает бренд Рядом и помогает возвращаться к совместным планам.')
pic(nativeplate,16.5)
h('Единая терминология')
p('Веб и приложение используют «Встречи», «Мои встречи», «Профиль» и одинаковые названия действий. Слова «событие» и «мероприятие» допустимы в документации; в пользовательском интерфейсе основное понятие — «встреча».')
h('Ценность приложения')
p('Быстрый доступ к своим встречам, статусу заявок и актуальным условиям. Ссылка из приглашения или уведомления открывает нужную встречу. Создание, участие и управление сохраняют знакомый порядок действий, даже если навигация отличается от веба.')
h('Уведомления')
p('Предлагаемые поводы: решение по заявке, ответ на вопрос и изменение важных условий. Заголовок называет событие: «Заявка одобрена», «Изменилось место встречи». Текст предлагает конкретный шаг. Уведомление о новой встрече не должно создавать ложную срочность.')
h('Согласованность платформ')
p('Общие цвета и отступы берём из одной системы; начертания загружаем отдельно. Учитываем системный размер текста, safe area и видимость клавиатуры. Возврат из уведомления не должен терять контекст. Права на уведомления запрашиваем рядом с понятной пользой.')
caption('Это направление развития продукта. Состав функций приложения определяется командой; шесть основных сценариев сохраняют общий язык независимо от платформы.')


page(16,'Передача в разработку','Единые значения оформления позволяют сохранить характер бренда в React и будущем React Native приложении.')
h('Что включено в комплект')
p('Векторные знак, связка с названием, монохромная и обратная версии, квадратный мастер иконки. Значения оформления находятся в tokens.json; tokens.css и tokens.ts содержат те же цвета, интервалы и параметры для веба и приложения. states.json задаёт подписи состояний и визуальные роли; copy.ru.json содержит тексты сценариев. README объясняет использование.')
h('Как применять')
p('JSON — источник значений. В React веб-компоненты обращаются к CSS-переменным по смысловой роли: основной текст, фон, акцент, ошибка. В React Native используются числовые значения из TypeScript; шрифты и их начертания регистрируются отдельно. При изменении системы обновляются JSON, CSS и TypeScript.')
p('Общие компоненты проектируются как Button, Input, ActivityCard, CategoryChip и AvatarGroup. Поведение навигации, жестов и доступности реализуется с учётом платформы. Брендбук задаёт дизайн, а не подтверждает наличие этих функций.')
h('Проверка перед запуском')
p('1. Проверить кириллицу и загрузку всех начертаний.\n2. Проверить контраст реальных состояний, фокус и навигацию с клавиатуры.\n3. Проверить крупный системный текст, переносы и экран шириной 320 px.\n4. Проверить загрузку, пустую выдачу, ошибки, отмены и заполненную встречу.\n5. Подготовить платформенные иконки из квадратного мастера.\n6. Проверить название, знак и домен перед публичным запуском.')
h('Источники для реализации')
for s in [
'[1] W3C WCAG 2.2: контраст текста, использование цвета, фокус и размер целей. https://www.w3.org/TR/WCAG22/',
'[2] Montserrat: официальный репозиторий, кириллица и лицензия SIL OFL 1.1. https://github.com/JulietaUla/Montserrat',
'[3] Open Sans: официальный репозиторий и лицензия. https://github.com/googlefonts/opensans',
'[4] Критерии оценки со слайда: шесть сценариев. Исходник сохранён в проекте: docs/references/evaluation-criteria.png.']:
    pp=p(s)
    for r in pp.runs:r.font.size=Pt(8.5)
caption('Версия 1.1 — основа для дизайна продукта. Проверка доступности выполняется на работающих экранах; готовое соответствие стандарту этим документом не заявляется.')

# The bundled Word template has a title border. Remove it explicitly.
for root in (doc.styles.element,doc.element):
    for el in list(root.iter(qn('w:pBdr'))): el.getparent().remove(el)
doc.save(OUT/'ryadom-brandbook.docx')
with zipfile.ZipFile(OUT/'ryadom-brand-assets.zip','w',zipfile.ZIP_DEFLATED) as z:
    for path in sorted(AS.iterdir()):z.write(path,'ryadom/'+path.name)
print(OUT/'ryadom-brandbook.docx')
print(OUT/'ryadom-brand-assets.zip')
